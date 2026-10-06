#!/usr/bin/env python3
"""Quarterly 4g check-in numbers (/2026-4g-checkin), first run 2026-10-06 for Q3.

For a calendar quarter (default: the most recently completed one) it reports:

  1. All colors      days 0n!AG (⎣∀clr) is filled
  2. Basic habits    days 0n!AF (N color) is filled — AFTER the backfill below
  3. 1n done         average of 1n+!AN over the quarter's month-anchor weeks
  4. 2n+             Σ 1n+ row 88 over the quarter's three months / 4500
  5. Kids outdoor    Σ 1n+!AL (1 kids nature) over the quarter's weeks
  6. hcmc -1         days 0n!AV (minutes tagged #-1) > 0, and total minutes

Backfill (JM's rule, 2026-10-06): a day with 0t and 0l both filled but AF
blank finished the basic habits; AF gets 3000. Written through the
excel-http daemon (ledger src "4g backfill"), re-checked cell by cell first.

Reads a fresh scp of the live workbook from ix (never the Straylight mirror).

Usage: 4g_checkin.py [YYYY-Qn] [--dry-run]
"""
from __future__ import annotations

import datetime as dt
import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci

REPO = Path.home() / "i446-monorepo"
COLS = json.loads((REPO / "config" / "neon-cols.json").read_text())["sheets"]
TWO_N_TARGET = 4500
BACKFILL_VALUE = "3000"

_spec = importlib.util.spec_from_file_location("gneon", REPO / "tools/4gneon-s/4gneon-s.py")
gneon = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gneon)


def col(sheet: str, header: str) -> int:
    """0-based index of a header's column, from neon-cols.json."""
    return ci(COLS[sheet]["headers"][header]) - 1


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    return v if isinstance(v, dt.date) else None


def blank(v) -> bool:
    return v is None or v == ""


def quarter_days(ws0n, start, end):
    """[(row_number, date, row_values)] for elapsed days in the quarter."""
    out = []
    today = dt.date.today()
    for i, r in enumerate(ws0n.iter_rows(min_row=2, values_only=True), start=2):
        d = as_date(r[2] if len(r) > 2 else None)
        if d and start <= d <= min(end, today):
            out.append((i, d, list(r)))
    return out


def backfill(days, dry_run: bool) -> list:
    """Write 3000 into AF where 0t and 0l are filled but AF is blank."""
    q0t, q0l, af = col("0n", "0t"), col("0n", "0l"), col("0n", "N color")
    get = lambda r, i: r[i] if i < len(r) else None
    cands = [(row, d) for row, d, r in days
             if not blank(get(r, q0t)) and not blank(get(r, q0l)) and blank(get(r, af))]
    if dry_run or not cands:
        return cands
    sys.path.insert(0, str(REPO / "lib"))
    from neon import excel
    done = []
    for row, d in cands:
        if excel.read("0n", "AF", row=row).get("formula"):
            continue  # filled since the snapshot: leave it
        if excel.write("0n", "AF", row=row, value=BACKFILL_VALUE,
                       src="4g backfill: 0t+0l done, AF blank -> 3000").get("ok"):
            done.append((row, d))
    return done


def main(argv: list[str]) -> int:
    dry = "--dry-run" in argv
    args = [a for a in argv if not a.startswith("--")]
    year, q, start, end = gneon.resolve_quarter(args[0] if args else None)
    months = [start.month, start.month + 1, start.month + 2]

    wb = openpyxl.load_workbook(gneon.fetch_live_workbook(), data_only=True, read_only=True)
    ws0n, ws1n = wb["0n"], wb["1n+"]
    days = quarter_days(ws0n, start, end)
    n = len(days)
    pct = lambda k: f"{k}/{n} days ({k / n * 100:.0f}%)" if n else "no days"
    get = lambda r, i: r[i] if i < len(r) else None

    ag, af, av = col("0n", "⎣∀clr"), col("0n", "N color"), col("0n", "-1")
    all_colors = sum(1 for _, _, r in days if not blank(get(r, ag)) and get(r, ag) != 0)
    filled = backfill(days, dry)
    basic = sum(1 for _, _, r in days if not blank(get(r, af))) + len(filled)
    hcmc = [get(r, av) for _, _, r in days]
    hcmc_days = sum(1 for v in hcmc if isinstance(v, (int, float)) and v > 0)
    hcmc_min = sum(v for v in hcmc if isinstance(v, (int, float)) and v > 0)

    # 1n+ weekly rows: col B is the M.W label (float artifacts like 7.199999).
    an, al = col("1n+", "avg"), col("1n+", "1 kids nature")
    an_vals, kids = [], 0
    for r in ws1n.iter_rows(min_row=6, max_row=60, values_only=True):
        wk = r[1] if len(r) > 1 else None
        if not isinstance(wk, (int, float)) or int(round(wk, 1)) not in months:
            continue
        if isinstance(get(r, an), (int, float)):
            an_vals.append((round(wk, 1), get(r, an)))
        if isinstance(get(r, al), (int, float)):
            kids += get(r, al)

    # 2n: row 61 holds month numbers 1-12; row 88 is that month's 2n points.
    # The table has no year column, so refuse when the sheet is for another year.
    sheet_year = next(ws0n.iter_rows(min_row=1, max_row=1, values_only=True))[2]
    r61 = next(ws1n.iter_rows(min_row=61, max_row=61, values_only=True))
    r88 = next(ws1n.iter_rows(min_row=88, max_row=88, values_only=True))
    two_n = {m: r88[i] for i, m in enumerate(r61) if m in months and i < len(r88)}
    two_n_sum = sum(v for v in two_n.values() if isinstance(v, (int, float)))

    print(f"4g check-in — {year} Q{q} ({start} to {end}), {n} days elapsed"
          + ("  [dry run: no writes]" if dry else ""))
    print(f"1. All colors (AG):      {pct(all_colors)}")
    verb = "would fill" if dry else "filled"
    print(f"2. Basic habits (AF):    {pct(basic)}  ({verb} {len(filled)} blank day(s) with 3000)")
    if filled:
        print("   " + ", ".join(str(d)[5:] for _, d in filled))
    if an_vals:
        avg = sum(v for _, v in an_vals) / len(an_vals) * 100
        print(f"3. 1n done:              {avg:.1f}%  [" + ", ".join(f"{w}={v * 100:.0f}%" for w, v in an_vals) + "]")
    else:
        print("3. 1n done:              no AN data for this quarter's weeks")
    if isinstance(sheet_year, (int, float)) and int(sheet_year) != year:
        print(f"4. 2n+:                  skipped — row 88 is built for {int(sheet_year)}, not {year}")
    else:
        detail = ", ".join(f"m{m}={two_n.get(m)}" for m in months)
        print(f"4. 2n+:                  {two_n_sum:g} / {TWO_N_TARGET} = {two_n_sum / TWO_N_TARGET * 100:.1f}%  [{detail}]")
    print(f"5. Kids outdoor (AL):    {kids:g} points")
    print(f"6. hcmc -1 (AV):         {pct(hcmc_days)} with #-1 minutes, {hcmc_min:g} min total")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
