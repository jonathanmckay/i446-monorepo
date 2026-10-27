#!/usr/bin/env python3
"""classify: file loose documents into the ~/files domain tree with Haiku.

(2026-10-27: retargeted from Dropbox to ~/files, home-folder-tree.md; areas grow as
Dropbox folders are imported. label.py runs the same catalog interactively.)

The destination list is read live from ~/files (config/triage.json "classify"):
every folder under the configured areas, cut to `depth` levels with years
templated (`qz12 - Finance/Income Taxes/{year} Tax`), that holds at least
`min_files` files. Each destination is shown to the model with a few sample
filenames, so it files the way JM already did. Renaming folders needs no code
change: the next run reads the new names.

Per file the model sees the name plus the first ~1500 chars of text (pdftotext,
textutil, openpyxl, pptx XML). Dropbox is almost entirely online-only and
reading a file downloads it, so content is only read below `max_hydrate_mb`;
--names-only reads nothing.

  eval      hide where already-filed files live, let the model place them,
            score exact-destination and area accuracy per confidence threshold.
            Leave-one-out: a test file never appears among the sample names.
  propose   dry run over a folder's top-level files; TSV to stdout.
  apply     propose, then move every file at or above `threshold`; logged to
            ~/.local/state/jm/triage/YYYY-MM-DD.jsonl like triage, so undoable.

Calls go through `claude -p --model haiku` on the subscription login (hooks off,
ANTHROPIC_API_KEY stripped so nothing bills the API account), `batch` files per
call, a few calls in parallel.

Usage: classify.py eval [--n 300] [--seed 1] [--names-only] [--model sonnet]
       classify.py propose|apply <dir> [--names-only]
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import os
import random
import re
import subprocess
import sys
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
CONFIG = REPO / "config" / "triage.json"
STATE = Path.home() / ".local" / "state" / "jm" / "triage"
YEAR_RE = re.compile(r"(?<!\d)((?:19|20)\d\d)(?!\d)")
SF_DATALESS = 0x40000000          # online-only File Provider placeholder
SNIPPET = 1500
# Subscription login, not the pay-per-token key in ~/.zshrc: with the key set (or
# under --bare, which skips the keychain) `claude -p` bills the API account.
CLAUDE_ENV = {k: v for k, v in os.environ.items() if k != "ANTHROPIC_API_KEY"}


def load_config() -> dict:
    cfg = json.loads(CONFIG.read_text())["classify"]
    cfg["root"] = Path(cfg["root"]).expanduser()
    return cfg


# ── destinations ────────────────────────────────────────────────────────────

@dataclass
class Dest:
    key: str                                   # templated, relative to root
    files: list[Path] = field(default_factory=list)


YEARLY_RE = re.compile(r"\D*(?:19|20)\d\d\D*")


def _templated(parent: Path, name: str) -> str:
    """`2019 Tax` -> `{year} Tax` when the year is the name's only number and a
    sibling follows the same pattern; one-off dated folders keep their name."""
    if not YEARLY_RE.fullmatch(name):
        return name
    t = YEAR_RE.sub("{year}", name)
    sibs = sum(1 for s in parent.iterdir() if s.is_dir() and YEARLY_RE.fullmatch(s.name)
               and YEAR_RE.sub("{year}", s.name) == t)
    return t if sibs >= 2 else name


def dest_key(root: Path, rel_dir: str, depth: int) -> str:
    parts, cur = [], root
    for c in rel_dir.split("/")[:depth]:
        parts.append(_templated(cur, c))
        cur = cur / c
    return "/".join(parts)


def catalog(cfg: dict) -> dict[str, Dest]:
    root, skip = cfg["root"], cfg["skip"]
    out: dict[str, Dest] = {}
    for area in cfg["areas"]:
        if not (root / area).is_dir():
            continue
        for d in sorted((root / area).rglob("*")):
            rel = str(d.relative_to(root))
            if (not d.is_dir() or any(rel == s or rel.startswith(s + "/") for s in skip)
                    or "_files" in rel or "/." in "/" + rel):
                continue
            files = [f for f in d.iterdir() if f.is_file() and not f.name.startswith(".")]
            if files:
                key = dest_key(root, rel, cfg["depth"])
                out.setdefault(key, Dest(key)).files += files
        top = root / area                       # files sitting directly in the area
        files = [f for f in top.iterdir() if f.is_file() and not f.name.startswith(".")]
        if files:
            out.setdefault(area, Dest(area)).files += files
    return {k: v for k, v in out.items() if len(v.files) >= cfg["min_files"]}


def resolve(cfg: dict, key: str, year: str | None) -> Path | None:
    """Concrete folder for a templated key; None when {year} can't be filled."""
    if "{year}" in key:
        if not year:
            return None
        root = cfg["root"]
        parts = key.split("/")
        cur = root
        for part in parts:          # reuse an existing folder whose name templates to this part
            hits = [c for c in (cur.iterdir() if cur.is_dir() else [])
                    if c.is_dir() and YEAR_RE.sub("{year}", c.name) == part
                    and ("{year}" not in part or year in c.name)]
            cur = hits[0] if hits else cur / part.replace("{year}", year)
        return cur
    return cfg["root"] / key


# ── content ─────────────────────────────────────────────────────────────────

def _run(cmd: list[str]) -> str:
    try:
        return subprocess.run(cmd, capture_output=True, text=True, timeout=30).stdout
    except (OSError, subprocess.TimeoutExpired):
        return ""


def snippet(p: Path, max_mb: float) -> str:
    try:
        st = p.stat()
    except OSError:
        return ""
    if st.st_flags & SF_DATALESS and st.st_size > max_mb * 1e6:
        return ""
    ext = p.suffix.lower()
    text = ""
    try:
        if ext == ".pdf":
            text = _run(["pdftotext", "-l", "2", "-q", str(p), "-"])
        elif ext in (".doc", ".docx", ".rtf", ".odt", ".html", ".htm", ".webarchive"):
            text = _run(["textutil", "-convert", "txt", "-stdout", str(p)])
        elif ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            bits = ["sheets: " + ", ".join(wb.sheetnames)]
            for row in wb[wb.sheetnames[0]].iter_rows(max_row=15, values_only=True):
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    bits.append(" | ".join(cells)[:200])
            wb.close()
            text = "\n".join(bits)
        elif ext == ".pptx":
            with zipfile.ZipFile(p) as z:
                slides = sorted(n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))[:3]
                text = " ".join(re.sub(r"<[^>]+>", " ", z.read(n).decode("utf8", "ignore")) for n in slides)
        elif ext in (".txt", ".md", ".csv", ".json", ".eml"):
            with p.open("r", errors="ignore") as f:
                text = f.read(SNIPPET * 2)
    except Exception:               # corrupt / encrypted / odd formats: name only
        return ""
    text = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", text)
    return re.sub(r"\s+", " ", text).strip()[:SNIPPET]


# ── model ───────────────────────────────────────────────────────────────────

SYSTEM = """You file JM's documents into his ~/files folder tree. Folder names use his domain codes \
(qz12 = personal finance, m5c7/m5x2 = McKay Capital real estate funds and properties, h335 = career, \
g245 = goals, i447 = infrastructure/tech, xk88 = marriage/social, s897 = social, hcb = health, \
o314 = journal, m828 = nonprofit, f8 = career fast track). {year} in a folder means one folder per year.

Destinations (number: path | sample files already there):
{dests}

For each file you are given, pick the single best destination number, or null if none fits. \
If the destination contains {{year}}, give the 4-digit year the document belongs to (tax year, \
statement year), else null. confidence is your probability (0-1) that a careful human would file \
it in exactly that destination.

Reply with only a JSON array, one object per file, no prose:
[{{"id": 0, "dest": 12, "year": "2019", "confidence": 0.85}}]"""


def system_prompt(dests: list[Dest], hide: set[Path], rng: random.Random, samples: int) -> str:
    lines = []
    for i, d in enumerate(dests):
        pool = [f.name for f in d.files if f not in hide]
        names = rng.sample(pool, min(samples, len(pool)))
        lines.append(f"{i}: {d.key} | " + "; ".join(n[:70] for n in names))
    return SYSTEM.replace("{dests}", "\n".join(lines))


def _call(system: str, batch: list[tuple[int, str, str]], model: str) -> list[dict]:
    user = "\n\n".join(f"id {i}\nname: {name}\ncontent: {text or '(not read)'}" for i, name, text in batch)
    for _ in range(2):
        raw = _run_claude(["claude", "-p", "--model", model, "--output-format", "json",
                           "--settings", '{"disableAllHooks": true}', "--system-prompt", system], user)
        try:
            result = json.loads(raw)["result"]
            arr = json.loads(result[result.index("["): result.rindex("]") + 1])
            return [r for r in arr if isinstance(r, dict) and "id" in r]
        except (ValueError, KeyError, TypeError):
            continue
    print(f"  model reply unreadable for ids {[b[0] for b in batch]}", file=sys.stderr)
    return []


def _run_claude(cmd: list[str], prompt: str) -> str:
    try:
        return subprocess.run(cmd, input=prompt.replace("\0", " "), capture_output=True, text=True,
                              timeout=300, cwd="/tmp", env=CLAUDE_ENV).stdout
    except (OSError, subprocess.TimeoutExpired):
        return ""


@dataclass
class Guess:
    path: Path
    key: str | None
    year: str | None
    confidence: float


def classify(paths: list[Path], dests: list[Dest], cfg: dict, *, names_only: bool,
             hide: set[Path] = frozenset(), seed: int = 0) -> list[Guess]:
    rng = random.Random(seed)
    system = system_prompt(dests, set(hide), rng, cfg["samples"])
    with ThreadPoolExecutor(8) as ex:
        texts = list(ex.map(lambda p: "" if names_only else snippet(p, cfg["max_hydrate_mb"]), paths))
    items = [(i, p.name, t) for i, (p, t) in enumerate(zip(paths, texts))]
    batches = [items[i:i + cfg["batch"]] for i in range(0, len(items), cfg["batch"])]
    by_id: dict[int, dict] = {}
    with ThreadPoolExecutor(cfg["parallel"]) as ex:
        for rows in ex.map(lambda b: _call(system, b, cfg["model"]), batches):
            for r in rows:
                by_id[int(r["id"])] = r
    out = []
    for i, p in enumerate(paths):
        r = by_id.get(i, {})
        n = r.get("dest")
        key = dests[n].key if isinstance(n, int) and 0 <= n < len(dests) else None
        year = str(r["year"]) if r.get("year") else (m.group(1) if (m := YEAR_RE.search(p.name)) else None)
        out.append(Guess(p, key, year, float(r.get("confidence") or 0)))
    return out


# ── commands ────────────────────────────────────────────────────────────────

def area_of(key: str) -> str:
    parts = key.split("/")
    return "/".join(parts[:2]) if parts[0] == "Archive" else parts[0]


def cmd_eval(cfg: dict, n: int, seed: int, names_only: bool) -> dict:
    cat = catalog(cfg)
    dests = list(cat.values())
    rng = random.Random(seed)
    per = max(1, n // len(dests) + 1)       # spread across destinations, not file-weighted
    picks: list[tuple[Path, str]] = []
    for d in dests:
        if len(d.files) > cfg["min_files"]:  # leaves samples to show after hiding
            picks += [(f, d.key) for f in rng.sample(d.files, min(per, len(d.files) - cfg["min_files"]))]
    rng.shuffle(picks)
    picks = picks[:n]
    truth = {p: k for p, k in picks}
    guesses = classify([p for p, _ in picks], dests, cfg, names_only=names_only,
                       hide=set(truth), seed=seed)

    rows, curve = [], []
    for g in guesses:
        rows.append({"file": str(g.path.relative_to(cfg["root"])), "truth": truth[g.path],
                     "guess": g.key, "confidence": g.confidence,
                     "exact": g.key == truth[g.path],
                     "area": g.key is not None and area_of(g.key) == area_of(truth[g.path])})
    for t in (0.0, 0.5, 0.7, 0.8, 0.9, 0.95):
        kept = [r for r in rows if r["guess"] and r["confidence"] >= t]
        curve.append({"threshold": t, "coverage": len(kept) / max(1, len(rows)),
                      "exact": sum(r["exact"] for r in kept) / max(1, len(kept)),
                      "area": sum(r["area"] for r in kept) / max(1, len(kept))})
    confusions = collections.Counter((area_of(r["truth"]), area_of(r["guess"]))
                                     for r in rows if r["guess"] and not r["area"])
    report = {"run": dt.datetime.now().isoformat(timespec="seconds"), "model": cfg["model"], "n": len(rows),
              "destinations": len(dests), "names_only": names_only, "curve": curve,
              "area_confusions": [[a, b, c] for (a, b), c in confusions.most_common(10)],
              "rows": rows}
    STATE.mkdir(parents=True, exist_ok=True)
    out = STATE / f"classify-eval-{dt.date.today()}-{cfg['model']}{'-names' if names_only else ''}.json"
    out.write_text(json.dumps(report, indent=1, ensure_ascii=False))
    report["path"] = str(out)
    return report


def cmd_propose(cfg: dict, folder: Path, names_only: bool) -> list[tuple[Guess, Path | None]]:
    dests = list(catalog(cfg).values())
    files = sorted(p for p in folder.iterdir() if p.is_file() and not p.name.startswith("."))
    return [(g, resolve(cfg, g.key, g.year) if g.key else None)
            for g in classify(files, dests, cfg, names_only=names_only)]


def cmd_apply(cfg: dict, proposals: list[tuple[Guess, Path | None]]) -> int:
    sys.path.insert(0, str(Path(__file__).parent))
    import triage
    log = STATE / f"{dt.date.today()}.jsonl"
    done = 0
    with log.open("a") as f:
        for g, dest in proposals:
            if dest is None or g.confidence < cfg["threshold"]:
                continue
            action = "trash" if triage.same_bytes(g.path, dest / g.path.name) else "move"
            dst = triage._move(g.path, triage.TRASH if action == "trash" else dest)
            f.write(json.dumps({"ts": dt.datetime.now().isoformat(timespec="seconds"),
                                "rule": "classify", "action": action, "confidence": g.confidence,
                                "src": str(g.path), "dst": str(dst)}, ensure_ascii=False) + "\n")
            done += 1
    return done


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("command", choices=["eval", "propose", "apply"])
    ap.add_argument("folder", nargs="?")
    ap.add_argument("--n", type=int, default=300)
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--names-only", action="store_true")
    ap.add_argument("--model", help="override config model (haiku, sonnet)")
    a = ap.parse_args()
    cfg = load_config()
    if a.model:
        cfg["model"] = a.model
    if a.command == "eval":
        r = cmd_eval(cfg, a.n, a.seed, a.names_only)
        print(f"{r['n']} files, {r['destinations']} destinations{' (names only)' if a.names_only else ''}")
        print("threshold  coverage  exact  area")
        for c in r["curve"]:
            print(f"  {c['threshold']:.2f}     {c['coverage']:6.0%}  {c['exact']:5.0%}  {c['area']:5.0%}")
        for a_, b, c in r["area_confusions"]:
            print(f"  confused {a_} -> {b}: {c}")
        print(r["path"])
        return
    if not a.folder:
        ap.error("propose/apply need a folder")
    proposals = cmd_propose(cfg, Path(a.folder).expanduser(), a.names_only)
    for g, dest in proposals:
        rel = str(dest.relative_to(cfg["root"])) if dest else "-"
        print(f"{g.confidence:.2f}\t{g.path.name}\t{rel}")
    if a.command == "apply":
        print(f"moved {cmd_apply(cfg, proposals)} at confidence >= {cfg['threshold']}", file=sys.stderr)


if __name__ == "__main__":
    main()
