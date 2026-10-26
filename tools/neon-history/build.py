#!/usr/bin/env python3
"""neon-history: one queryable SQLite DB of every Neon workbook version.

WHAT
    Reads every known Neon 分 workbook (2018 Quip exports through the live
    Neon分v12.2.xlsx), copies each to a scratch dir first (originals are never
    opened by Excel or written), and loads them into

        ~/vault/z_asts/g245--neon-history--neon-history.sqlite

    The live workbook is pulled read-only from Ix with scp (the Straylight
    OneDrive mirror can lag days); --no-ix uses the local mirror instead.
    Re-running rebuilds the DB from scratch (idempotent). Takes ~3-5 minutes,
    dominated by openpyxl parsing the 8-30 MB 2024-2026 workbooks twice
    (once for cached values, once for formulas).

USAGE
    python3 ~/i446-monorepo/tools/neon-history/build.py            # build
    python3 ~/i446-monorepo/tools/neon-history/build.py --list     # show sources
    sqlite3 ~/vault/z_asts/g245--neon-history--neon-history.sqlite \
        "select * from daily where date >= '2026-03-01' limit 5"

TABLES
    sources(source_version, path, era, role, first_date, last_date, n_days, note)
        role = canonical (owns its dates in the `daily` view) or snapshot
        (an older copy of the same year; kept for drift checks, core metrics only).
    obs(date, source_version, metric, value, derived)
        The long table. `derived` = 1 when the cell was a formula that
        references other cells (e.g. the end-of-day remainder `=D-SUM(...)`),
        0 for typed numbers or literal arithmetic like `=3+150`.
    canon(date, source_version)    which version owns each date.
    domain_map(header, grp)        raw 0分 column header -> comparable group.

VIEWS
    canon_obs   obs restricted to the canonical version for each date.
    toggl_daily Toggl per-day cross-check (wake time, night sleep, 04-10 minutes
                by project group); source_version='toggl' in obs, 2018 onward.
    daily       one row per date: total, b1..b9, morning, blocks_ok,
                n_time, allclr_time, sleep_min, habits_done, habits_avail.

METRICS (modern era, 2023-10 onward; 0分 / 0₹ sheet + 0n / 0₦ sheet)
    total            Σ column (sum of the in-day domain columns).
    b1..b9           points entered per block. 2023-10..2026-04 the blocks
                     were the Arabic prayer-time names; JM's own 2026 header
                     migration mapped them by position to 地支:
                       b1 فجر=卯  b2 شروق=辰  b3 صباح=巳  b4 ظهر=午  b5 عصر=未
                       b6 آصيل=申  b7 غروب=酉  b8 غسق=戌  b9 صدفة/亥 (remainder)
                     Arabic era blocks were odd-hour (05-07, ...), 地支 even-hour
                     (04-06, ...). b9 absorbs whatever was not attributed.
    bflag1..bflag8   the Arabic-era per-block marker cell (usually 1).
    morning          b1+b2+b3 (only meaningful when blocks_ok = 1).
    blocks_ok        1 when b1..b3 are typed values (not a formula lump) and at
                     least 3 of b1..b8 are typed numbers > 1, i.e. the day was
                     actually attributed to blocks rather than dumped in one cell.
    n_time / allclr_time   0₦t / ∀₦t: HHMM when basic habits / all colors were
                     done (0 or blank = not done; 3000 = backfilled marker).
    sleep_min        0n 睡觉 minutes.
    habits_done / habits_avail   0n columns with a row-2 weight that hold a
                     value > 0, over the number of such columns that year.
    col:<header>     every other 0分 header column, raw.
    hab:<header>     every 0n habit column, raw (minutes or marks).
    raw:<sheet>:<header>   Quip-era (2017-2023Q3) cells; dates inferred from
                     row order and validated against the day-of-week column.
                     Those sheets have no comparable daily total and their
                     block cells are 0/1 flags or fractions, not points.

Neon off Excel note: the live ledger (~/vault/g245/neon-ledger, lib/neon/ledger.py)
is the per-write record from 2026-07 on; this DB is the per-day history.
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import openpyxl  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HOME = Path.home()
OD = HOME / "OneDrive/vault-excel"
ODP = HOME / "Library/CloudStorage/OneDrive-Personal/vault-excel"
ARCH = ODP / "z_archive"
OLD = ARCH / "Old + Documents"
QUIP = ARCH / "Old Neon (2017-2023)"
DB = Path(os.environ.get("NEON_HISTORY_DB", HOME / "vault/z_asts/g245--neon-history--neon-history.sqlite"))

# (source_version, candidate paths, era, role, note). First existing path wins.
SOURCES = [
    ("v12.2-live", [OD / "Neon分v12.2.xlsx", ODP / "Neon分v12.2.xlsx"], "modern", "canonical",
     "live workbook, 2026"),
    ("v12.1", [OLD / "Neon分v12.1.xlsx"], "modern", "snapshot", "2026 copy saved 2026-03-20"),
    ("v12.3", [OLD / "Neon分v12.3.xlsx", HOME / "Downloads/Neon分v12.3.xlsx"], "modern", "snapshot",
     "2026 copy saved 2026-03-27"),
    ("v11", [OLD / "Neon分v11.xlsx"], "modern", "snapshot", "2026 copy saved 2026-03-04"),
    ("v10", [OLD / "Neon Canonical Archive/Neon分v10.xlsx"], "modern", "canonical", "2025"),
    ("v9", [OLD / "Neon分v9.xlsx"], "modern", "snapshot", "2025 copy saved 2025-05-24"),
    ("v81", [OLD / "Neon Canonical Archive/Neon分v81.xlsx"], "modern", "canonical", "2024"),
    ("v8", [OLD / "Neon分v8.xlsx"], "modern", "snapshot", "2024 copy saved 2024-12-21"),
    ("v7.151", [OLD / "2024.Q1 霓虹系统 ₦ ^LLM 分 v7.151 (g245) (Neon).xlsx"], "modern", "snapshot",
     "2024 Q1"),
    ("v7.0", [OLD / "2023.Q4 霓虹系统 ₦ + 分 v7 (g245) (Neon).xlsx"], "modern", "canonical", "2023 Q4"),
]
_QUIP = [
    ("q2018", "quip-2017-2018-霓虹系统-neon-v1-g245.xlsx"),
    ("q2019q1", "quip-2019q1-霓虹系统-neon-v3-g245.xlsx"),
    ("q2019q2", "quip-2019q2-霓虹系统-neon-v4-g245.xlsx"),
    ("q2019q3", "quip-2019q3-霓虹系统-neon-v41-g245.xlsx"),
    ("q2019q4", "quip-2019q4-霓虹系统-neon-v42-g245.xlsx"),
    ("q2020q1", "quip-2020q1-霓虹系统-neon-v43-g245.xlsx"),
    ("q2020q2", "quip-2020q2-霓虹系统-neon-v50-g245.xlsx"),
    ("q2020q3", "quip-2020q3-霓虹系统-v51-g245-neon-safe.xlsx"),
    ("q2020q4", "quip-2020q4-霓虹系统-v511-g245-neon.xlsx"),
    ("q2021q1", "quip-2021q1-霓虹系统-v52-g245-neon.xlsx"),
    ("q2021q2", "quip-2021q2-霓虹系统-v53-g245-neon.xlsx"),
    ("q2021q3", "quip-2021q3-霓虹系统-v54-g245-neon.xlsx"),
    ("q2021q4", "quip-2021q4-霓虹系统-v60-g245-neon.xlsx"),
    ("q2022q1", "quip-2022q1-霓虹系统-v61-g245-neon.xlsx"),
    ("q2022q2", "quip-2022q2-霓虹系统-v62-g245-neon.xlsx"),
    ("q2022q3", "quip-2022q3-霓虹系统-v63-g245-neon.xlsx"),
    ("q2022q4", "quip-2022q4-霓虹系统-v64-g245-neon.xlsx"),
    ("q2023q1", "quip-2023q1-霓虹系统-v65-g245-neon.xlsx"),
    ("q2023q2", "quip-2023q2-霓虹系统-v66-g245-neon.xlsx"),
    ("q2023q3", "quip-2023q3-霓虹系统-v67-g245-neon.xlsx"),
]
SOURCES += [(v, [QUIP / f], "quip", "canonical", "Quip export, values only") for v, f in _QUIP]

QUIP_ORDER = {v: i for i, (v, _) in enumerate(_QUIP)}
# Canonical precedence when two canonical sources cover the same date (lower wins).
PRECEDENCE = {"v12.2-live": 0, "v10": 1, "v81": 2, "v7.0": 3}

ARABIC = ["فجر", "شروق", "صباح", "ظهر", "عصر", "آصيل", "غروب", "غسق", "صدفة"]
DIZHI = ["卯", "辰", "巳", "午", "未", "申", "酉", "戌", "亥"]
POINT_SHEETS = ("0分", "0₹")
HABIT_SHEETS = ("0n", "0₦")
CORE = {"total", "n_time", "allclr_time", "sleep_min", "habits_done", "habits_avail"} | {
    f"b{i}" for i in range(1, 10)}
REF_RE = re.compile(r"(?<![A-Za-z])\$?[A-Z]{1,3}\$?\d+")

# Raw 0分 header -> comparable group. Habit points moved from a dedicated 0₦
# column (2023-2025) into the domain columns in 2026, so `habits` vs the
# domain groups is NOT comparable across that boundary; see the vault note.
DOMAIN_MAP = {
    "-1₦": "ritual", "0₦": "habits", "1₦": "weekly", "0₲": "goals", "0g": "goals", "-2₲": "goals",
    "i8": "work_i9", "i9": "work_i9", "mp": "work_i9", "m7": "m5x2", "m5": "m5x2",
    "媒体消": "media", "媒体": "media", "媒": "media", "思": "mind",
    "吃": "health", "动": "health", "hcb": "health", "个": "self",
    "f6": "people_f6", "c84": "honor_c84", "家": "family", "xk": "family", "社": "social",
    "d": "other",
}


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            return float(s)
    return None


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def is_derived(f) -> int:
    return 1 if isinstance(f, str) and f.startswith("=") and REF_RE.search(f[1:]) else 0


def resolve(paths):
    for p in paths:
        if Path(p).exists():
            return Path(p)
    return None


def load_pair(path: Path, sheets):
    """Return {sheet: (value_rows, formula_rows)} for the first matching sheet names."""
    out = {}
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    wf = openpyxl.load_workbook(path, read_only=True, data_only=False)
    for name in sheets:
        if name in wv.sheetnames:
            mr = min(wv[name].max_row or 0, 1200)
            mc = min(wv[name].max_column or 0, 120)
            vals = [list(r) for r in wv[name].iter_rows(max_row=mr, max_col=mc, values_only=True)]
            fms = [list(r) for r in wf[name].iter_rows(max_row=mr, max_col=mc, values_only=True)]
            out[name] = (vals, fms)
    wv.close()
    wf.close()
    return out


def cell(rows, i, j):
    return rows[i][j] if i < len(rows) and j < len(rows[i]) else None


# ---------------------------------------------------------------- modern era
def parse_points(vals, fms):
    """0分/0₹ sheet. Yields (date, metric, value, derived)."""
    # header row: the row whose column B says 5^0
    hr = next((i for i in range(4) if str(cell(vals, i, 1) or "").strip() == "5^0"), None)
    if hr is None:
        return
    hdrs = {}
    for i in range(hr, min(hr + 2, len(vals))):
        for j, v in enumerate(vals[i]):
            if isinstance(v, str) and v.strip() and j not in hdrs:
                hdrs[j] = v.strip()
    blocks, flags, special, domains = {}, {}, {}, {}
    for j, h in hdrs.items():
        if h in ARABIC:
            k = ARABIC.index(h) + 1
            if k not in blocks:
                blocks[k] = j + 1
                if k <= 8:
                    flags[k] = j
        elif h in DIZHI:
            blocks.setdefault(DIZHI.index(h) + 1, j)
        elif h == "Σ":
            special.setdefault("total", j)
        elif h in ("0₦t",):
            special.setdefault("n_time", j)
        elif h in ("∀₦t", "%∀c"):
            special.setdefault("allclr_time", j)
        elif j > 2 and h not in ("周", "天", "地", "5^0") and not h.isdigit():
            domains.setdefault(j, h)
    for i in range(hr + 1, len(vals)):
        d = to_date(cell(vals, i, 1))
        if not d:
            continue
        for m, j in special.items():
            x = num(cell(vals, i, j))
            if x is not None:
                yield d, m, x, is_derived(cell(fms, i, j))
        for k, j in blocks.items():
            x = num(cell(vals, i, j))
            if x is not None:
                yield d, f"b{k}", x, is_derived(cell(fms, i, j))
        for k, j in flags.items():
            x = num(cell(vals, i, j))
            if x is not None:
                yield d, f"bflag{k}", x, is_derived(cell(fms, i, j))
        for j, h in domains.items():
            x = num(cell(vals, i, j))
            if x is not None:
                yield d, f"col:{h}", x, is_derived(cell(fms, i, j))


SKIP_HAB = re.compile(r"^(∑|⎣|₦∑|N color|#|-?\d+\+?$|2\+|1\+|蠢之后|睡觉$|cpap$)")


def parse_habits(vals, fms):
    """0n/0₦ sheet: row 1 names, row 2 weights (⊖分), dates in col B or C."""
    dc = max(range(4), key=lambda j: sum(1 for r in vals if j < len(r) and to_date(r[j])))
    names = {j: str(v).strip() for j, v in enumerate(vals[0]) if v not in (None, "") and j > dc}
    weights = {j: num(cell(vals, 1, j)) for j in names}
    habit_cols = [j for j, n in names.items()
                  if weights.get(j) and not SKIP_HAB.match(n)]
    special = {}
    for j, n in names.items():
        if n == "睡觉":
            special.setdefault("sleep_min", j)
        elif n in ("⎣₦", "N color"):
            special.setdefault("n_time_0n", j)
        elif n == "⎣∀clr":
            special.setdefault("allclr_time_0n", j)
    for i in range(2, len(vals)):
        d = to_date(cell(vals, i, dc))
        if not d:
            continue
        for m, j in special.items():
            x = num(cell(vals, i, j))
            if x is not None:
                yield d, m, x, is_derived(cell(fms, i, j))
        done = 0
        any_val = False
        for j in habit_cols:
            x = num(cell(vals, i, j))
            if x is not None:
                any_val = True
                yield d, f"hab:{names[j]}", x, is_derived(cell(fms, i, j))
                if x > 0:
                    done += 1
        if any_val:
            yield d, "habits_done", float(done), 0
            yield d, "habits_avail", float(len(habit_cols)), 0


# ---------------------------------------------------------------- quip era
DOW = {"S": None, "Su": 6, "M": 0, "T": None, "Tu": 1, "W": 2, "Th": 3, "F": 4, "Sa": 5}


def parse_quip(path: Path, version: str, log):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(max_row=min(ws.max_row or 0, 400),
                                             max_col=min(ws.max_column or 0, 80), values_only=True)]
        if len(rows) < 30:
            continue
        # Daily sheets: column B holds day-of-week letters on data rows.
        dow_rows = [i for i, r in enumerate(rows) if len(r) > 1 and r[1] in DOW]
        anchors = [(i, to_date(r[0])) for i, r in enumerate(rows) if r and to_date(r[0])]
        if len(dow_rows) < 40 or not anchors:
            continue
        # Quip lost the date formulas; usually only the first row keeps a date,
        # and some sheets carry stale dates copied from an earlier quarter.
        # Accept the first anchor that puts the first data row within 10 days
        # of the file's quarter start (version names encode it: q2021q3).
        m = re.match(r"q(\d{4})q(\d)", version)
        qstart = dt.date(int(m.group(1)), 3 * int(m.group(2)) - 2, 1) if m else None
        pick = None
        for a_i, a_d in anchors:
            first = a_d + dt.timedelta(days=dow_rows[0] - a_i)
            if qstart is None or abs((first - qstart).days) <= 10:
                pick = (a_i, a_d)
                break
        if pick is None:
            log(f"  skip {version}/{ws.title}: no date anchor near {qstart}")
            continue
        a_i, a_d = pick
        # validate offset arithmetic against the DOW letters
        ok = tot = 0
        for i in dow_rows:
            d = a_d + dt.timedelta(days=i - a_i)
            want = DOW[rows[i][1]]
            letter = rows[i][1]
            if want is None:  # ambiguous S / T
                want_set = {5, 6} if letter == "S" else {1, 3}
            else:
                want_set = {want}
            tot += 1
            ok += d.weekday() in want_set
        if tot == 0 or ok / tot < 0.9:
            log(f"  skip {version}/{ws.title}: weekday check {ok}/{tot}")
            continue
        hdr = {}
        for r in rows[:3]:
            for j, v in enumerate(r):
                if isinstance(v, str) and v.strip() and j > 1:
                    hdr[j] = (hdr[j] + "|" if j in hdr else "") + v.strip()[:16]
        n = 0
        for i in dow_rows:
            d = a_d + dt.timedelta(days=i - a_i)
            for j, h in hdr.items():
                x = num(rows[i][j]) if j < len(rows[i]) else None
                if x is not None:
                    n += 1
                    yield d, f"raw:{ws.title}:{h}", x, 0
        log(f"  {version}/{ws.title}: {len(dow_rows)} days, weekday check {ok}/{tot}, {n} cells")
    wb.close()

# ---------------------------------------------------------------- Toggl
TOGGL_DB = HOME / "vault/i156/toggl.sqlite"
TOGGL_GROUPS = {
    "work": {"i9", "m5x2", "infra", "i444", "i447", "g245", "epcn", "n156", "qz12", "h335",
             "m828", "q5n7", "f8"},
    "family": {"xk87", "xk88", "家", "s897"},
    "health": {"hcb", "hcbp", "hci"},
    "mind": {"hcm"},
    "media": {"hcmc", "hcmc2"},
    "sleep": {"睡觉"},
}
MORNING = (4, 10)  # 卯+辰+巳


def load_toggl(con, log):
    """Per-day Toggl aggregates as source_version='toggl' (independent of Neon's
    own bookkeeping, so usable as a cross-check across eras):
      tg_am_<group>_min   minutes 04:00-10:00 by project group
      tg_am_tracked_min   all tracked minutes 04:00-10:00
      tg_wake             HHMM of the last 睡觉 entry ending 03:00-12:00
      tg_night_sleep_min  睡觉 minutes between D-1 18:00 and D 12:00
    Wall-clock times are taken as recorded (travel time zones not normalised)."""
    if not TOGGL_DB.exists():
        log("toggl: missing")
        return
    src = sqlite3.connect(f"file:{TOGGL_DB}?mode=ro", uri=True)
    grp = {p: g for g, ps in TOGGL_GROUPS.items() for p in ps}
    agg = {}
    for start, stop, project in src.execute(
            "SELECT start, stop, project FROM entries WHERE stop IS NOT NULL"):
        try:
            a = dt.datetime.fromisoformat(start[:19])
            b = dt.datetime.fromisoformat(stop[:19])
        except ValueError:
            continue
        if b <= a or (b - a).days > 1:
            continue
        g = grp.get(project, "other")
        day = a.date()
        while day <= b.date():
            w0 = dt.datetime.combine(day, dt.time(MORNING[0]))
            w1 = dt.datetime.combine(day, dt.time(MORNING[1]))
            ov = (min(b, w1) - max(a, w0)).total_seconds() / 60
            if ov > 0:
                k = day.isoformat()
                agg[(k, f"tg_am_{g}_min")] = agg.get((k, f"tg_am_{g}_min"), 0) + ov
                agg[(k, "tg_am_tracked_min")] = agg.get((k, "tg_am_tracked_min"), 0) + ov
            if g == "sleep" and b.date() == day and 3 <= b.hour < 12:
                    k = (day.isoformat(), "tg_wake")
                    agg[k] = max(agg.get(k, 0), b.hour * 100 + b.minute)
            day += dt.timedelta(days=1)
        if g == "sleep":  # night ending on day D = D-1 18:00 .. D 12:00
            for day in {a.date(), a.date() + dt.timedelta(days=1)}:
                n0 = dt.datetime.combine(day - dt.timedelta(days=1), dt.time(18))
                n1 = dt.datetime.combine(day, dt.time(12))
                ov2 = (min(b, n1) - max(a, n0)).total_seconds() / 60
                if ov2 > 0:
                    k = (day.isoformat(), "tg_night_sleep_min")
                    agg[k] = agg.get(k, 0) + ov2
    src.close()
    con.executemany("INSERT OR REPLACE INTO obs VALUES (?,?,?,?,0)",
                    [(d, "toggl", m, v) for (d, m), v in agg.items()])
    ds = sorted({d for d, _ in agg})
    con.execute("INSERT INTO sources VALUES (?,?,?,?,?,?,?,?)",
                ("toggl", str(TOGGL_DB), "toggl", "aux", ds[0], ds[-1], len(ds),
                 "Toggl time entries, aggregated per day"))
    log(f"toggl: {len(agg)} obs, {len(ds)} days {ds[0]}..{ds[-1]}")


# ---------------------------------------------------------------- build
SCHEMA = """
CREATE TABLE sources(source_version TEXT PRIMARY KEY, path TEXT, era TEXT, role TEXT,
  first_date TEXT, last_date TEXT, n_days INTEGER, note TEXT);
CREATE TABLE obs(date TEXT NOT NULL, source_version TEXT NOT NULL, metric TEXT NOT NULL,
  value REAL, derived INTEGER NOT NULL DEFAULT 0,
  PRIMARY KEY(date, source_version, metric)) WITHOUT ROWID;
CREATE TABLE canon(date TEXT PRIMARY KEY, source_version TEXT NOT NULL);
CREATE TABLE domain_map(header TEXT PRIMARY KEY, grp TEXT NOT NULL);
CREATE VIEW canon_obs AS
  SELECT o.* FROM obs o JOIN canon c ON c.date = o.date AND c.source_version = o.source_version;
CREATE VIEW block_quality AS
  SELECT date,
    SUM(CASE WHEN metric IN ('b1','b2','b3') AND derived = 1 THEN 1 ELSE 0 END) AS morning_derived,
    SUM(CASE WHEN metric IN ('b1','b2','b3','b4','b5','b6','b7','b8')
             AND derived = 0 AND value > 1 THEN 1 ELSE 0 END) AS typed_blocks
  FROM canon_obs GROUP BY date;
CREATE VIEW toggl_daily AS
  SELECT date,
    MAX(CASE WHEN metric='tg_wake' THEN value END) AS wake_hhmm,
    MAX(CASE WHEN metric='tg_night_sleep_min' THEN value END) AS night_sleep_min,
    MAX(CASE WHEN metric='tg_am_tracked_min' THEN value END) AS am_tracked_min,
    MAX(CASE WHEN metric='tg_am_work_min' THEN value END) AS am_work_min,
    MAX(CASE WHEN metric='tg_am_family_min' THEN value END) AS am_family_min,
    MAX(CASE WHEN metric='tg_am_sleep_min' THEN value END) AS am_sleep_min,
    MAX(CASE WHEN metric='tg_am_mind_min' THEN value END) AS am_mind_min,
    MAX(CASE WHEN metric='tg_am_health_min' THEN value END) AS am_health_min,
    MAX(CASE WHEN metric='tg_am_media_min' THEN value END) AS am_media_min
  FROM obs WHERE source_version = 'toggl' GROUP BY date;
CREATE VIEW daily AS
  SELECT c.date, c.source_version,
    MAX(CASE WHEN metric='total' THEN value END) AS total,
    MAX(CASE WHEN metric='b1' THEN value END) AS b1,
    MAX(CASE WHEN metric='b2' THEN value END) AS b2,
    MAX(CASE WHEN metric='b3' THEN value END) AS b3,
    MAX(CASE WHEN metric='b4' THEN value END) AS b4,
    MAX(CASE WHEN metric='b5' THEN value END) AS b5,
    MAX(CASE WHEN metric='b6' THEN value END) AS b6,
    MAX(CASE WHEN metric='b7' THEN value END) AS b7,
    MAX(CASE WHEN metric='b8' THEN value END) AS b8,
    MAX(CASE WHEN metric='b9' THEN value END) AS b9,
    COALESCE(MAX(CASE WHEN metric='b1' THEN value END),0)
      + COALESCE(MAX(CASE WHEN metric='b2' THEN value END),0)
      + COALESCE(MAX(CASE WHEN metric='b3' THEN value END),0) AS morning,
    CASE WHEN q.morning_derived = 0 AND q.typed_blocks >= 3 THEN 1 ELSE 0 END AS blocks_ok,
    MAX(CASE WHEN metric='n_time' THEN value END) AS n_time,
    MAX(CASE WHEN metric='allclr_time' THEN value END) AS allclr_time,
    MAX(CASE WHEN metric='sleep_min' THEN value END) AS sleep_min,
    MAX(CASE WHEN metric='habits_done' THEN value END) AS habits_done,
    MAX(CASE WHEN metric='habits_avail' THEN value END) AS habits_avail
  FROM canon c
  JOIN canon_obs o ON o.date = c.date
  LEFT JOIN block_quality q ON q.date = c.date
  GROUP BY c.date;
"""


def fetch_live(tmpdir: Path, log):
    """The OneDrive mirror on non-Ix Macs can lag days (see memory note on
    OneDrive stale sync). Ix is canonical, so pull a read-only copy from it
    when we are not on Ix; fall back to the local mirror if ssh fails."""
    if "ix" in os.uname().nodename.lower().split(".")[0].split("-"):
        return None
    dst = tmpdir / "live-from-ix.xlsx"
    r = subprocess.run(["scp", "-q", "-o", "ConnectTimeout=10",
                        "ix:OneDrive/vault-excel/Neon分v12.2.xlsx", str(dst)],
                       capture_output=True, timeout=180)
    if r.returncode == 0 and dst.exists():
        log("live workbook: fresh copy from Ix")
        return dst
    log("live workbook: ssh ix failed, using local OneDrive mirror (may be stale)")
    return None


def build(verbose=True, live: Path | None = None, from_ix: bool = True):
    log = print if verbose else (lambda *a: None)
    tmpdir = Path(tempfile.mkdtemp(prefix="neon-history-"))
    if live is None and from_ix:
        live = fetch_live(tmpdir, log)
    DB.parent.mkdir(parents=True, exist_ok=True)
    tmp_db = DB.with_suffix(".sqlite.tmp")
    if tmp_db.exists():
        tmp_db.unlink()
    con = sqlite3.connect(tmp_db)
    con.executescript(SCHEMA)
    con.executemany("INSERT INTO domain_map VALUES (?,?)", DOMAIN_MAP.items())
    try:
        for version, paths, era, role, note in SOURCES:
            src = live if (version == "v12.2-live" and live) else resolve(paths)
            if not src:
                log(f"missing: {version}")
                continue
            local = tmpdir / f"{version}.xlsx"
            shutil.copy2(src, local)  # read the copy, never the original
            log(f"{version}: {src.name}")
            rows = []
            if era == "modern":
                sheets = load_pair(local, POINT_SHEETS + HABIT_SHEETS)
                for name, (v, f) in sheets.items():
                    gen = parse_points(v, f) if name in POINT_SHEETS else parse_habits(v, f)
                    for d, m, x, der in gen:
                        if role == "snapshot" and m not in CORE:
                            continue
                        rows.append((d.isoformat(), version, m, x, der))
            else:
                for d, m, x, der in parse_quip(local, version, log):
                    rows.append((d.isoformat(), version, m, x, der))
            today = dt.date.today().isoformat()
            rows = [r for r in rows if r[0] <= today]  # drop empty template rows ahead
            con.executemany("INSERT OR REPLACE INTO obs VALUES (?,?,?,?,?)", rows)
            ds = sorted({r[0] for r in rows})
            con.execute("INSERT INTO sources VALUES (?,?,?,?,?,?,?,?)",
                        (version, "ix:OneDrive/vault-excel/Neon分v12.2.xlsx" if src == live else str(src), era, role, ds[0] if ds else None,
                         ds[-1] if ds else None, len(ds), note))
            log(f"  {len(rows)} obs, {len(ds)} days {ds[0] if ds else ''}..{ds[-1] if ds else ''}")
            local.unlink()
        # canonical owner per date: modern by PRECEDENCE, then quip (later file wins)
        canon = {}
        for (d, v) in con.execute(
                "SELECT DISTINCT o.date, o.source_version FROM obs o JOIN sources s "
                "USING(source_version) WHERE s.role='canonical'"):
            rank = PRECEDENCE.get(v, 100 - QUIP_ORDER.get(v, 0))  # later Quip file wins
            if d not in canon or rank < canon[d][1]:
                canon[d] = (v, rank)
        load_toggl(con, log)
        con.executemany("INSERT INTO canon VALUES (?,?)",
                        [(d, v) for d, (v, _) in sorted(canon.items())])
        con.commit()
        con.execute("VACUUM")
    finally:
        con.close()
        shutil.rmtree(tmpdir, ignore_errors=True)
    os.replace(tmp_db, DB)
    log(f"wrote {DB} ({DB.stat().st_size / 1e6:.1f} MB)")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--list", action="store_true", help="list sources and exit")
    ap.add_argument("-q", action="store_true", help="quiet")
    ap.add_argument("--live", type=Path, help="path to a copy of the live workbook to use")
    ap.add_argument("--no-ix", action="store_true", help="don't scp the live workbook from Ix")
    a = ap.parse_args()
    if a.list:
        for version, paths, era, role, note in SOURCES:
            p = resolve(paths)
            print(f"{version:12} {role:9} {era:6} {'OK ' if p else 'MISSING'} {p or paths[0]}")
        return
    build(verbose=not a.q, live=a.live, from_ix=not a.no_ix)


if __name__ == "__main__":
    sys.exit(main())
