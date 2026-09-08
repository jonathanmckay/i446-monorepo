#!/usr/bin/env python3
"""refresh_points_cache — build the points JSON cache from Neon Excel.

Extracted from refresh-points-cache.sh's embedded `python3 -c "..."` block
(2026-09-08): that form is nearly untestable and, worse, an unescaped `"`
inside a Python comment silently truncated the whole script wherever bash's
outer double-quoting parsed it — no error, no output, no file write, no
alerting (found and fixed same day). A real .py file gets real error
handling and a real regression test instead of a shell-quoting trap.

Uses openpyxl (data_only=True) which reads cached formula values from the
last Excel save — NOT xlwings/live Excel, since this scans up to 90 days of
history per sheet and a live-Excel per-cell read (dashboard.py's approach)
would be far too slow for that. The tradeoff: cron-spawned processes on
macOS have intermittently failed to even open the OneDrive-hosted .xlsx via
plain file I/O with `PermissionError: Operation not permitted` — a Full
Disk Access (TCC) restriction that doesn't apply to an interactive SSH
session's process tree, only to some spawners (confirmed live 2026-09-08:
identical manual runs succeeded every time; cron's own log
(/tmp/refresh-points-cache.log) showed repeated PermissionErrors over the
same period). That's a one-time macOS Full Disk Access grant for cron/bash
in System Settings, not something this script can fix at the Python level —
but it CAN fail loudly instead of a silent no-op, which is what actually
made today's bug take so long to find.
"""
from __future__ import annotations

import json
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci

NEON = Path.home() / "OneDrive/vault-excel/Neon分v12.2.xlsx"
CACHE = Path(__file__).parent / ".points-cache.json"
CUTOFF_DAYS = 90

# 0分: per-domain points + block breakdown + day grand total (col D).
# __total__ is read by neg1n's /api/day-points (quarter-circle arc complication).
FEN_COLS = {16: "-1₦", 17: "0₲", 18: "i9", 19: "m5", 20: "个", 21: "媒",
            22: "思", 23: "hcb", 24: "xk", 25: "社"}
BLOCK_COLS = {7: "卯", 8: "辰", 9: "巳", 10: "午", 11: "未", 12: "申",
              13: "酉", 14: "戌", 15: "亥"}

# hcbi: calories eaten (col U) + today's hcbp+hcbc score (Y+AA), per day,
# for neg1n's /api/hcb. hcbp+hcbc was originally read as the fixed Q2+Q3
# running total (hcbi!X375+X378) but that's a year-scale figure incompatible
# with the 131 goal, which is a DAILY target — JM (2026-09-08):
# "=hcbi!AA{row}+hcbi!Y{row}" is the actual per-day figure to use.
KCAL_COL = ci("U")
HCBP_HCBC_COLS = (ci("Y"), ci("AA"))

# 0n: prayer count (صلاة) + hcmp minutes (o314+冥想+其他人), per day, for
# neg1n's /api/hcmp.
SALAT_COL = ci("AP")
HCMP_COLS = (ci("AQ"), ci("AR"), ci("AS"))


def as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def load_workbook_or_die(path: Path):
    try:
        return openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except PermissionError as e:
        sys.exit(
            f"ERROR: can't open {path} ({e}). This is almost always macOS Full "
            "Disk Access: the process running this script needs it granted in "
            "System Settings -> Privacy & Security -> Full Disk Access (an "
            "interactive SSH/Terminal session usually already has it; cron's "
            "own spawned process may not). Not retrying — a permission wall "
            "doesn't clear itself."
        )
    except Exception as e:
        sys.exit(f"ERROR: can't open {path}: {type(e).__name__}: {e}")


def build_fen_days(ws, today: date, cutoff: date, result: dict) -> None:
    for row in ws.iter_rows(min_row=3, values_only=True):
        d = as_date(row[1])
        if d is None or d <= cutoff or d > today:
            continue
        day_data = result.setdefault(d.isoformat(), {})
        for idx, label in FEN_COLS.items():
            val = row[idx - 1]
            if val is not None and isinstance(val, (int, float)) and val > 0:
                day_data[label] = int(round(float(val)))
        block_data = {}
        for idx, label in BLOCK_COLS.items():
            val = row[idx - 1]
            if val is not None and isinstance(val, (int, float)) and val > 0:
                block_data[label] = int(round(float(val)))
        if block_data:
            day_data["__block__"] = block_data
        total = row[ci("D") - 1]
        if isinstance(total, (int, float)):
            day_data["__total__"] = int(round(float(total)))


def build_hcbi_days(ws, today: date, cutoff: date, result: dict) -> None:
    for row in ws.iter_rows(min_row=3, values_only=True):
        d = as_date(row[1])
        if d is None or d <= cutoff or d > today:
            continue
        day_data = result.setdefault(d.isoformat(), {})
        kcal = row[KCAL_COL - 1]
        if isinstance(kcal, (int, float)) and kcal > 0:
            day_data["__hcb_kcal__"] = int(round(float(kcal)))
        hcbp_hcbc = sum(row[c - 1] for c in HCBP_HCBC_COLS if isinstance(row[c - 1], (int, float)))
        day_data["__hcbp_hcbc__"] = int(round(float(hcbp_hcbc)))


def build_0n_days(ws, today: date, cutoff: date, result: dict) -> None:
    for row in ws.iter_rows(min_row=5, values_only=True):
        d = as_date(row[2])
        if d is None or d <= cutoff or d > today:
            continue
        day_data = result.setdefault(d.isoformat(), {})
        salat = row[SALAT_COL - 1]
        if isinstance(salat, (int, float)):
            day_data["__salat__"] = int(round(float(salat)))
        hcmp_min = sum(row[c - 1] for c in HCMP_COLS if isinstance(row[c - 1], (int, float)))
        day_data["__hcmp_min__"] = int(round(float(hcmp_min)))  # 0 is real, not missing


def build_cache(wb, today: date | None = None) -> dict:
    today = today or date.today()
    cutoff = today - timedelta(days=CUTOFF_DAYS)
    result: dict = {}
    build_fen_days(wb["0分"], today, cutoff, result)
    build_hcbi_days(wb["hcbi"], today, cutoff, result)
    build_0n_days(wb["0n"], today, cutoff, result)
    return result


def main() -> int:
    wb = load_workbook_or_die(NEON)
    try:
        result = build_cache(wb)
    finally:
        wb.close()
    CACHE.write_text(json.dumps(result, indent=2, ensure_ascii=False))
    print(f"wrote {len(result)} days to cache")
    return 0


if __name__ == "__main__":
    sys.exit(main())
