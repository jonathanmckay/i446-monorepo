#!/usr/bin/env python3
"""3s.py — /3s: fill a quarter's row of the scorecard's 3s tab from Neon.

Quarters are WEEK quarters (4-4-5): 13 Sunday-anchored weeks, label-months
3q-2..3q of the 1n+ fiscal ladder (Q3 2026 = weeks 7.1..9.5 = Sun 7/5 – Sat
10/3), via lib/neon/weeks.py. Not calendar quarters.

Columns written in scorecard.xlsx › '16-26 3s' (row: col A == 'YYYY.Qn'):
  B  Timestamp   now
  D  0₲✓         days with any 0g in 0分!Q (> 0)         → '=n/days'
  E  sd 0₦       days with 0 < 0分!E (0₦t) < 1000        → '=n/days'
  F  com0        days with 0分!E nonzero                 → '=n/days'
  G  1₦          mean of 1n+!AN on the quarter's 3 month-anchor rows (M.1);
                 refuses if any is blank
  H  2₦          1n+ row 89 under the quarter's first month (row 89 is a
                 rolling 3-month window: Σ row 88 over 3 months / (D88·3));
                 blank → 0
  K  m5c7 Rating mean of i9+m5x2!G (m5x2) over the quarter's weeks
  L  I9 Rating   mean of i9+m5x2!B (i9) over the quarter's weeks
                 ratings: MM=1 MA=2 EE=3, OL/blank skipped; mean → nearest
                 letter, '+'/'-' when >= .25 off it
  AB hcb sum    mean per day of 0分!W (hcb: hcbi AA + Y + direct appends)
  AC hcbp       mean per day of hcbi!Y
  AD hcbc 分    mean per day of hcbi!AA (∑c)
  I (3₦) and every other column are left alone.

Also renames the tab '16-23 3s' → '16-26 3s' if the old name is still there.

Neon is read from ix's file (scp; never the Straylight OneDrive mirror, which
desyncs). The scorecard is written through LOCAL Excel: it is open only on
Straylight (checked 2026-10-05).

Usage:
  3s.py [YYYY-Qn | Qn] [--partial] [--dry-run]
  No quarter → the most recently COMPLETED week quarter. An in-progress
  quarter is refused unless --partial (then percentages are over elapsed days).
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci

sys.path.insert(0, str(Path.home() / "i446-monorepo" / "lib"))
from neon import weeks  # noqa: E402

REMOTE_NEON = "~/OneDrive/vault-excel/Neon分v12.2.xlsx"
SCORECARD = "scorecard.xlsx"
SCORECARD_PATH = Path.home() / "OneDrive/vault-excel/scorecard.xlsx"
SHEET = "16-26 3s"
OLD_SHEETS = ("16-23 3s",)
RATING = {"MM": 1, "MA": 2, "EE": 3}
LETTER = {1: "MM", 2: "MA", 3: "EE"}


# --- pure helpers (unit-tested) ---------------------------------------------
def quarter_weeks(year: int, q: int) -> list[date]:
    """The 13 Sundays whose fiscal label-month is in quarter q."""
    months = {3 * q - 2, 3 * q - 1, 3 * q}
    d = weeks.first_sunday_of_year(year)
    out = []
    while len(out) < 13 and d.year <= year + 1:
        m = int(weeks.fiscal_week_label(d).split(".")[0])
        if m in months:
            out.append(d)
        elif out:
            break
        d += timedelta(days=7)
    if len(out) != 13:
        raise ValueError(f"{year} Q{q}: found {len(out)} weeks, expected 13")
    return out


def resolve_quarter(arg: str | None, today: date) -> tuple[int, int]:
    if arg:
        m = re.match(r"^(?:(\d{4})-)?Q([1-4])$", arg.strip(), re.IGNORECASE)
        if not m:
            raise SystemExit(f"bad quarter {arg!r}; use 2026-Q3 or Q3")
        return int(m.group(1) or today.year), int(m.group(2))
    label_month = int(weeks.fiscal_week_label(weeks.week_sunday(today)).split(".")[0])
    q = (label_month - 1) // 3 + 1 - 1
    return (today.year - 1, 4) if q == 0 else (today.year, q)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def rating_mean(vals: list) -> tuple[str | None, int]:
    scores = [RATING[str(v).strip().upper()] for v in vals
              if v is not None and str(v).strip().upper() in RATING]
    if not scores:
        return None, 0
    avg = sum(scores) / len(scores)
    base = min(3, max(1, round(avg)))
    off = avg - base
    sfx = "+" if off >= 0.25 else "-" if off <= -0.25 else ""
    return LETTER[base] + sfx, len(scores)


def label_of(v) -> str:
    return f"{v:.1f}" if isinstance(v, float) else str(v).strip()


# --- Neon read ---------------------------------------------------------------
def fetch_neon() -> Path:
    tmp = Path(tempfile.mkdtemp()) / "neon.xlsx"
    subprocess.run(["scp", "-q", f"ix:{REMOTE_NEON}", str(tmp)], check=True,
                   capture_output=True, text=True, timeout=120)
    return tmp


def compute(wb, year: int, q: int, today: date, partial: bool) -> dict:
    sundays = quarter_weeks(year, q)
    start, end = sundays[0], sundays[-1] + timedelta(days=6)
    labels = [weeks.fiscal_week_label(s) for s in sundays]
    if end >= today and not partial:
        raise SystemExit(f"{year} Q{q} ({start}..{end}) is not over; pass --partial to score it so far")
    last = min(end, today - timedelta(days=1)) if partial else end

    days = []
    for r in wb["0分"].iter_rows(min_row=2, max_row=600, values_only=True):
        d = r[1]
        if not isinstance(d, (date, datetime)):
            continue
        d = d.date() if isinstance(d, datetime) else d
        if start <= d <= last:
            days.append({"date": d, "E": num(r[4]), "Q": num(r[16]), "W": num(r[22])})
    n = len(days)
    expected = (last - start).days + 1
    if n != expected:
        raise SystemExit(f"0分 has {n} rows for {start}..{last}, expected {expected}")
    g0 = sum(1 for d in days if (d["Q"] or 0) > 0)   # any 0g that day
    sd = sum(1 for d in days if d["E"] and 0 < d["E"] < 1000)
    com = sum(1 for d in days if d["E"])
    hcb = sum(d["W"] or 0 for d in days) / n

    # hcbi is keyed by its own date column (rows are offset from 0分's).
    hy, ha = {}, {}
    for r in wb["hcbi"].iter_rows(min_row=2, max_row=600, values_only=True):
        d = r[1]
        if isinstance(d, (date, datetime)):
            d = d.date() if isinstance(d, datetime) else d
            if start <= d <= last:
                hy[d], ha[d] = num(r[ci("Y") - 1]), num(r[ci("AA") - 1])
    if len(hy) != n:
        raise SystemExit(f"hcbi has {len(hy)} rows for {start}..{last}, expected {n}")
    hcbp = sum(v or 0 for v in hy.values()) / n
    hcbc = sum(v or 0 for v in ha.values()) / n

    ws1 = wb["1n+"]
    an_col = ci("AN") - 1
    anchors = {f"{m}.1" for m in (3 * q - 2, 3 * q - 1, 3 * q)}
    an = {}
    for r in ws1.iter_rows(min_row=4, max_row=70, values_only=True):
        lab = label_of(r[1])
        if lab in anchors:
            an[lab] = num(r[an_col])
    missing = sorted(a for a in anchors if an.get(a) is None)
    if missing and not partial:
        raise SystemExit(f"1₦ invariant: 1n+!AN blank on anchor week(s) {missing}; fill them first")
    an_vals = [v for v in an.values() if v is not None]
    one_n = sum(an_vals) / len(an_vals) if an_vals else None

    row61 = next(ws1.iter_rows(min_row=61, max_row=61, values_only=True))
    row89 = next(ws1.iter_rows(min_row=89, max_row=89, values_only=True))
    first_month = 3 * q - 2
    col = next((i for i, v in enumerate(row61) if num(v) == first_month), None)
    if col is None:
        raise SystemExit(f"2₦: no month-{first_month} column in 1n+ row 61")
    two_n = num(row89[col]) if col < len(row89) else None
    two_n_note = "" if two_n is not None else " (blank → 0)"
    two_n = two_n or 0.0

    i9, m5 = [], []
    for r in wb["i9+m5x2"].iter_rows(min_row=1, max_row=70, values_only=True):
        if label_of(r[0]) in labels:
            i9.append(r[1])
            m5.append(r[6] if len(r) > 6 else None)
    i9_r, i9_n = rating_mean(i9)
    m5_r, m5_n = rating_mean(m5)

    return {"year": year, "q": q, "start": start, "end": end, "last": last,
            "days": n, "g0": g0, "sd": sd, "com": com,
            "hcb": hcb, "hcbp": hcbp, "hcbc": hcbc,
            "one_n": one_n, "an": an, "two_n": two_n, "two_n_col": col,
            "two_n_note": two_n_note, "i9": i9_r, "i9_n": i9_n,
            "m5": m5_r, "m5_n": m5_n, "weeks": len(labels)}


# --- scorecard write -----------------------------------------------------------
def osa(script: str) -> str:
    r = subprocess.run(["osascript", "-"], input=script, capture_output=True,
                       text=True, timeout=120)
    if r.returncode != 0:
        raise SystemExit(f"Excel error: {r.stderr.strip()}")
    return r.stdout.strip()


def write(res: dict) -> str:
    label = f"{res['year']}.Q{res['q']}"
    n = res["days"]
    cells = {"D": f"={res['g0']}/{n}", "E": f"={res['sd']}/{n}",
             "F": f"={res['com']}/{n}", "H": f"{res['two_n']:.4f}",
             "AB": f"{res['hcb']:.1f}", "AC": f"{res['hcbp']:.1f}",
             "AD": f"{res['hcbc']:.1f}"}
    if res["one_n"] is not None:
        cells["G"] = f"{res['one_n']:.4f}"
    if res["m5"]:
        cells["K"] = res["m5"]
    if res["i9"]:
        cells["L"] = res["i9"]
    sets = "\n".join(
        f'    set formula of range ("{c}" & theRow) of ws to "{v}"' for c, v in cells.items())
    olds = " or ".join(f'n is "{o}"' for o in OLD_SHEETS)
    return osa(f'''tell application "Microsoft Excel"
    if (name of workbooks) does not contain "{SCORECARD}" then open POSIX file "{SCORECARD_PATH}"
    set wb to workbook "{SCORECARD}"
    set renamed to ""
    repeat with n in (get name of every worksheet of wb)
        set n to n as text
        if {olds} then
            set name of sheet n of wb to "{SHEET}"
            set renamed to " (renamed " & n & " → {SHEET})"
        end if
    end repeat
    set ws to sheet "{SHEET}" of wb
    set theRow to 0
    set lastRow to 1
    repeat with r from 2 to 200
        set a to (string value of range ("A" & r) of ws)
        if a is "{label}" then
            if theRow is not 0 then return "ERROR: {label} appears twice in col A"
            set theRow to r
        end if
        if a is not "" then set lastRow to r
    end repeat
    if theRow = 0 then
        set theRow to lastRow + 1
        set value of range ("A" & theRow) of ws to "{label}"
    end if
    set value of range ("B" & theRow) of ws to (current date)
{sets}
    save wb
    return "OK row " & theRow & renamed
end tell''')


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("quarter", nargs="?")
    ap.add_argument("--partial", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    today = date.today()
    year, q = resolve_quarter(a.quarter, today)
    path = fetch_neon()
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        res = compute(wb, year, q, today, a.partial)
    finally:
        path.unlink(missing_ok=True)
    n = res["days"]
    pct = lambda k: f"{res[k]}/{n} = {res[k] / n:.0%}"
    an = ", ".join(f"{k}={'—' if v is None else f'{v:.0%}'}" for k, v in sorted(res["an"].items()))
    print(f"3s {year}.Q{q}  weeks {weeks.fiscal_week_label(res['start'])}–"
          f"{weeks.fiscal_week_label(res['end'] - timedelta(days=6))}  "
          f"({res['start']} – {res['last']}, {n} days)")
    print(f"  0₲✓   {pct('g0')}   (0分!Q > 0: any 0g)")
    print(f"  sd 0₦ {pct('sd')}   (0 < 0分!E < 1000)")
    print(f"  com0  {pct('com')}   (0分!E logged)")
    one_n = "—" if res["one_n"] is None else f"{res['one_n']:.0%}"
    print(f"  1₦    {one_n}   (1n+!AN: {an})")
    print(f"  2₦    {res['two_n']:.0%}   (1n+ row 89, month {3 * q - 2}){res['two_n_note']}")
    print(f"  hcb   {res['hcb']:.1f}/day   (0分!W)   hcbp {res['hcbp']:.1f}/day (hcbi!Y)   "
          f"hcbc 分 {res['hcbc']:.1f}/day (hcbi!AA)")
    print(f"  m5c7  {res['m5'] or '—'}   ({res['m5_n']}/{res['weeks']} weeks rated)")
    print(f"  I9    {res['i9'] or '—'}   ({res['i9_n']}/{res['weeks']} weeks rated)")
    if a.dry_run:
        print("  dry run: nothing written")
        return 0
    out = write(res)
    if out.startswith("ERROR"):
        raise SystemExit(out)
    print(f"  → {SCORECARD} › {SHEET} {out[3:]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
